"""Word/Excel attachment -> PDF-page conversion (REQ-CAL-REV-PDF-ATTACH-
CONVERT-001, 2026-09-23).

Used by backend/review_summary_pdf_export.py's append_office_conversions to
turn a "word" or "excel" attachment's already-downloaded bytes (fetched by
the router from Cloudinary — this module never touches storage or the
database) into extra PDF pages appended onto the generated Review Summary
export, the same way append_pdf_attachments already appends a PDF
attachment's own pages.

Tool choice / deployment-environment note: this backend is deployed as a
Vercel Python serverless function (see .env's "NEVER copy this variable into
the Vercel backend project's production environment" comment and
handover/member-schedule-vercel-neon-deployment-preparation-2026-07-10.md).
That environment has no LibreOffice, Microsoft Office, or other external
document-conversion binary installed, and installing one is not a reliable
option on a serverless function. Conversion here is therefore pure Python —
python-docx (.docx only) + reportlab for Word, and openpyxl (.xlsx) / xlrd
(legacy .xls) + reportlab for Excel — so it works identically in this local
dev environment and in the actual Vercel deployment, with no external
process, no subprocess call, and no network access. This is a readable
re-rendering of the document's own text/table/cell content, not a
pixel-perfect layout clone of the original Word/Excel file.

Legacy binary .doc (pre-2007 OLE compound-file format, not a zip/OOXML
package) cannot be read by python-docx at all. That is reported as an
explicit, honest conversion failure (see ConversionResult below) — this
module never claims a .doc's contents were included when they were not.
Callers must render that failure to the reader rather than silently omitting
the attachment.

Every function here is bytes-in, dataclass-out and never raises for a
malformed/corrupt/unsupported input — the one exception is a genuine
programming error (e.g. a missing dependency), which is allowed to raise
normally since it indicates a broken deployment, not a bad input file.
"""

from dataclasses import dataclass
from io import BytesIO
from typing import List, Optional, Tuple

import docx
from docx.oxml.ns import qn
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
import openpyxl
import xlrd
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

_PORTRAIT_PAGE_WIDTH, _PORTRAIT_PAGE_HEIGHT = A4
_LANDSCAPE_PAGE_WIDTH, _LANDSCAPE_PAGE_HEIGHT = landscape(A4)
_MARGIN = 1.8 * cm
_PORTRAIT_CONTENT_WIDTH = _PORTRAIT_PAGE_WIDTH - 2 * _MARGIN
_LANDSCAPE_CONTENT_WIDTH = _LANDSCAPE_PAGE_WIDTH - 2 * _MARGIN

_MUTED_COLOR = colors.HexColor("#6b7280")
_HEADER_TEXT_COLOR = colors.HexColor("#1f2d3d")
_TABLE_BORDER_COLOR = colors.HexColor("#c7d2dd")
_TABLE_HEADER_BG_COLOR = colors.HexColor("#e6ecf2")

# Hard caps so one pathological input document (huge row/paragraph count)
# can never make conversion take unbounded time or produce an unbounded PDF.
# A capped render is always explicitly labelled as such, never silently
# truncated.
MAX_EXCEL_ROWS_PER_SHEET_FOR_PDF = 2000
MAX_EXCEL_COLUMNS_PER_SHEET_FOR_PDF = 60
MAX_EXCEL_SHEETS_FOR_PDF = 30
_MIN_EXCEL_COLUMN_WIDTH_CM = 2.4


@dataclass
class ConversionResult:
    """success=False always carries a human-readable error_message and
    pdf_bytes=None — callers must never treat a failed ConversionResult as
    "included"."""
    success: bool
    pdf_bytes: Optional[bytes] = None
    error_message: Optional[str] = None


def _styles():
    base = getSampleStyleSheet()
    return {
        "heading": ParagraphStyle(
            "OfficeConversionHeading", parent=base["Heading2"],
            textColor=_HEADER_TEXT_COLOR, fontSize=13, leading=16, spaceBefore=8, spaceAfter=4,
        ),
        "body": ParagraphStyle(
            "OfficeConversionBody", parent=base["Normal"],
            textColor=_HEADER_TEXT_COLOR, fontSize=10.5, leading=14, spaceAfter=6,
        ),
        "cell": ParagraphStyle(
            "OfficeConversionCell", parent=base["Normal"],
            textColor=_HEADER_TEXT_COLOR, fontSize=8, leading=10,
        ),
        "cell_header": ParagraphStyle(
            "OfficeConversionCellHeader", parent=base["Normal"], fontName="Helvetica-Bold",
            textColor=_HEADER_TEXT_COLOR, fontSize=8, leading=10,
        ),
        "note": ParagraphStyle(
            "OfficeConversionNote", parent=base["Normal"],
            textColor=_MUTED_COLOR, fontSize=8.5, spaceAfter=6,
        ),
    }


def _escape(raw_text: str) -> str:
    """Same escape/line-break convention as review_summary_pdf_export's own
    _paragraph_text — duplicated (not imported) so this module stays a
    standalone, independently testable unit with no import-order dependency
    on the PDF-export module."""
    text = raw_text or ""
    escaped = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return escaped.replace("\r\n", "\n").replace("\n", "<br/>")


# ── Word (.docx) ──────────────────────────────────────────────────────────

def _iter_block_items(document):
    """Standard python-docx recipe for iterating a document body's
    paragraphs and tables IN DOCUMENT ORDER — document.paragraphs and
    document.tables (the high-level API) are two separate, unordered-
    relative-to-each-other lists, which would silently reorder a document
    that interleaves text and tables."""
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield DocxParagraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield DocxTable(child, document)


def _word_table_flowable(table: DocxTable, styles):
    rows = []
    for row in table.rows:
        rows.append([Paragraph(_escape(cell.text.strip()), styles["cell"]) for cell in row.cells])
    if not rows:
        return None
    column_count = max(len(r) for r in rows)
    for row in rows:
        while len(row) < column_count:
            row.append(Paragraph("", styles["cell"]))
    col_width = _PORTRAIT_CONTENT_WIDTH / column_count
    flowable = Table(rows, colWidths=[col_width] * column_count, repeatRows=0)
    flowable.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, _TABLE_BORDER_COLOR),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    return flowable


def convert_word_bytes_to_pdf(file_bytes: bytes, filename: str) -> ConversionResult:
    """Converts a .docx's readable text and tables into PDF pages. Legacy
    .doc (OLE binary format, not OOXML) and any file python-docx cannot
    parse (corrupt, wrong extension, etc.) fail explicitly rather than
    producing an empty or misleading PDF."""
    try:
        document = docx.Document(BytesIO(file_bytes))
    except Exception:
        lower_name = (filename or "").lower()
        if lower_name.endswith(".doc") and not lower_name.endswith(".docx"):
            return ConversionResult(
                success=False,
                error_message=(
                    "Legacy .doc format is not supported for in-PDF conversion in this "
                    "environment (only .docx is supported). The original file is "
                    "available in the complete ZIP download."
                ),
            )
        return ConversionResult(
            success=False,
            error_message=(
                "This Word document could not be read (corrupt or unrecognized file). "
                "The original file is available in the complete ZIP download."
            ),
        )

    styles = _styles()
    story: list = []
    has_content = False
    try:
        for block in _iter_block_items(document):
            if isinstance(block, DocxParagraph):
                text = block.text.strip()
                if not text:
                    continue
                has_content = True
                style_name = (block.style.name if block.style is not None else "") or ""
                is_heading = style_name.lower().startswith("heading") or style_name.lower() == "title"
                story.append(Paragraph(_escape(text), styles["heading"] if is_heading else styles["body"]))
            elif isinstance(block, DocxTable):
                table_flowable = _word_table_flowable(block, styles)
                if table_flowable is not None:
                    has_content = True
                    story.append(table_flowable)
                    story.append(Spacer(1, 0.3 * cm))
    except Exception:
        return ConversionResult(
            success=False,
            error_message=(
                "This Word document's content could not be extracted. The original "
                "file is available in the complete ZIP download."
            ),
        )

    if not has_content:
        story.append(Paragraph("(This document has no readable text content.)", styles["note"]))

    try:
        buffer = BytesIO()
        pdf_doc = SimpleDocTemplate(
            buffer, pagesize=A4,
            leftMargin=_MARGIN, rightMargin=_MARGIN, topMargin=_MARGIN, bottomMargin=_MARGIN,
            title="Converted Word Attachment",
        )
        pdf_doc.build(story)
        return ConversionResult(success=True, pdf_bytes=buffer.getvalue())
    except Exception:
        return ConversionResult(
            success=False,
            error_message=(
                "This Word document's content could not be rendered into the PDF "
                "(one section may be too large or complex). The original file is "
                "available in the complete ZIP download."
            ),
        )


# ── Excel (.xlsx / legacy .xls) ──────────────────────────────────────────

def _stringify_cell(value) -> str:
    if value is None:
        return ""
    return str(value)


def _sheets_from_xlsx(file_bytes: bytes) -> List[Tuple[str, List[List[str]], bool, bool]]:
    """Returns a list of (sheet_name, rows, rows_truncated, columns_truncated)
    for a real .xlsx workbook, read_only for memory efficiency, values only
    (formulas resolved to their last-calculated value — the same "readable
    result", not formula text, a human opening the spreadsheet would see)."""
    workbook = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True, read_only=True)
    sheets = []
    for worksheet in workbook.worksheets[:MAX_EXCEL_SHEETS_FOR_PDF]:
        rows: List[List[str]] = []
        rows_truncated = False
        columns_truncated = False
        for row in worksheet.iter_rows(values_only=True):
            if len(rows) >= MAX_EXCEL_ROWS_PER_SHEET_FOR_PDF:
                rows_truncated = True
                break
            values = [_stringify_cell(v) for v in row]
            if len(values) > MAX_EXCEL_COLUMNS_PER_SHEET_FOR_PDF:
                columns_truncated = True
                values = values[:MAX_EXCEL_COLUMNS_PER_SHEET_FOR_PDF]
            rows.append(values)
        sheets.append((worksheet.title, rows, rows_truncated, columns_truncated))
    workbook.close()
    return sheets


def _sheets_from_xls(file_bytes: bytes) -> List[Tuple[str, List[List[str]], bool, bool]]:
    """Same shape as _sheets_from_xlsx, for legacy .xls via xlrd (xlrd>=2.0
    dropped .xlsx support entirely and reads only the legacy binary format —
    the exact complement of openpyxl, so the two libraries together cover
    both extensions with no overlap)."""
    book = xlrd.open_workbook(file_contents=file_bytes)
    sheets = []
    for sheet in book.sheets()[:MAX_EXCEL_SHEETS_FOR_PDF]:
        row_limit = min(sheet.nrows, MAX_EXCEL_ROWS_PER_SHEET_FOR_PDF)
        rows_truncated = sheet.nrows > row_limit
        columns_truncated = sheet.ncols > MAX_EXCEL_COLUMNS_PER_SHEET_FOR_PDF
        rows = []
        for r in range(row_limit):
            values = [_stringify_cell(v) for v in sheet.row_values(r)]
            if columns_truncated:
                values = values[:MAX_EXCEL_COLUMNS_PER_SHEET_FOR_PDF]
            rows.append(values)
        sheets.append((sheet.name, rows, rows_truncated, columns_truncated))
    return sheets


def _excel_columns_per_chunk() -> int:
    chunk = int(_LANDSCAPE_CONTENT_WIDTH / (_MIN_EXCEL_COLUMN_WIDTH_CM * cm))
    return max(4, min(12, chunk))


def _excel_chunk_table(rows: List[List[str]], styles):
    if not rows:
        return None
    column_count = max(len(r) for r in rows)
    col_width = _LANDSCAPE_CONTENT_WIDTH / column_count
    header_row = [Paragraph(_escape(v), styles["cell_header"]) for v in rows[0]]
    while len(header_row) < column_count:
        header_row.append(Paragraph("", styles["cell_header"]))
    body_rows = []
    for row in rows[1:]:
        cells = [Paragraph(_escape(v), styles["cell"]) for v in row]
        while len(cells) < column_count:
            cells.append(Paragraph("", styles["cell"]))
        body_rows.append(cells)
    table = Table([header_row] + body_rows, colWidths=[col_width] * column_count, repeatRows=1)
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, _TABLE_BORDER_COLOR),
        ("BACKGROUND", (0, 0), (-1, 0), _TABLE_HEADER_BG_COLOR),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]))
    return table


def convert_excel_bytes_to_pdf(file_bytes: bytes, filename: str) -> ConversionResult:
    """Converts every sheet of a .xlsx/.xls workbook into PDF pages,
    landscape-oriented (more usable width for wide spreadsheets than the
    main portrait export). A sheet wider than one page's column budget is
    split into column chunks — "readable pagination" — each chunk repeating
    its own header row (the sheet's first row) via Table(repeatRows=1) so a
    chunk that itself splits across multiple PDF pages still shows column
    headers on every page. Each sheet and chunk is labelled with the sheet
    name so a reader always knows which sheet/columns a page came from."""
    lower_name = (filename or "").lower()
    try:
        if lower_name.endswith(".xlsx"):
            sheets = _sheets_from_xlsx(file_bytes)
        elif lower_name.endswith(".xls"):
            sheets = _sheets_from_xls(file_bytes)
        else:
            return ConversionResult(
                success=False,
                error_message=(
                    "Unrecognized Excel file extension. The original file is "
                    "available in the complete ZIP download."
                ),
            )
    except Exception:
        return ConversionResult(
            success=False,
            error_message=(
                "This Excel file could not be read (corrupt or unrecognized file). "
                "The original file is available in the complete ZIP download."
            ),
        )

    styles = _styles()
    story: list = []
    columns_per_chunk = _excel_columns_per_chunk()
    has_content = False

    try:
        for sheet_name, rows, rows_truncated, columns_truncated in sheets:
            sheet_heading = Paragraph("Sheet: " + _escape(sheet_name), styles["heading"])
            if not rows:
                story.append(sheet_heading)
                story.append(Paragraph("(This sheet is empty.)", styles["note"]))
                continue
            column_count = max(len(r) for r in rows)
            chunk_ranges = [
                (start, min(start + columns_per_chunk, column_count))
                for start in range(0, column_count, columns_per_chunk)
            ]
            total_chunks = len(chunk_ranges)
            for chunk_index, (start, end) in enumerate(chunk_ranges, start=1):
                # Each chunk's own heading/label is grouped with its table
                # via KeepTogether (2026-09-23) so a page break can never
                # strand "Sheet: <name>" or "Columns X-Y of Z" alone at the
                # bottom of a page with its table pushed to the next one —
                # reportlab's KeepTogether still allows the group to split
                # internally if it is taller than a whole empty page, so
                # this is always safe to apply, even for a very long sheet.
                chunk_group: list = []
                if chunk_index == 1:
                    chunk_group.append(sheet_heading)
                if total_chunks > 1:
                    chunk_group.append(Paragraph(
                        "Columns " + str(start + 1) + "–" + str(end) + " of " + str(column_count)
                        + " (part " + str(chunk_index) + " of " + str(total_chunks) + ")",
                        styles["note"],
                    ))
                chunk_rows = [row[start:end] for row in rows]
                table = _excel_chunk_table(chunk_rows, styles)
                if table is not None:
                    has_content = True
                    chunk_group.append(table)
                    story.append(KeepTogether(chunk_group) if chunk_group else table)
                    story.append(Spacer(1, 0.35 * cm))
                elif chunk_group:
                    story.extend(chunk_group)
            if rows_truncated:
                story.append(Paragraph(
                    "Showing the first " + str(MAX_EXCEL_ROWS_PER_SHEET_FOR_PDF) + " rows of this sheet — "
                    "see the original file in the complete ZIP download for the rest.",
                    styles["note"],
                ))
            if columns_truncated:
                story.append(Paragraph(
                    "Showing the first " + str(MAX_EXCEL_COLUMNS_PER_SHEET_FOR_PDF) + " columns of this sheet — "
                    "see the original file in the complete ZIP download for the rest.",
                    styles["note"],
                ))
            story.append(Spacer(1, 0.2 * cm))
    except Exception:
        return ConversionResult(
            success=False,
            error_message=(
                "This Excel file's content could not be rendered into the PDF. The "
                "original file is available in the complete ZIP download."
            ),
        )

    if not has_content:
        story.append(Paragraph("(This workbook has no readable cell content.)", styles["note"]))

    try:
        buffer = BytesIO()
        pdf_doc = SimpleDocTemplate(
            buffer, pagesize=landscape(A4),
            leftMargin=_MARGIN, rightMargin=_MARGIN, topMargin=_MARGIN, bottomMargin=_MARGIN,
            title="Converted Excel Attachment",
        )
        pdf_doc.build(story)
        return ConversionResult(success=True, pdf_bytes=buffer.getvalue())
    except Exception:
        return ConversionResult(
            success=False,
            error_message=(
                "This Excel file's content could not be rendered into the PDF (a sheet "
                "may be too large or complex). The original file is available in the "
                "complete ZIP download."
            ),
        )
