"""Tests for backend/attachment_conversion.py (REQ-CAL-REV-PDF-ATTACH-
CONVERT-001, 2026-09-23) — Word/Excel attachment -> PDF-page conversion.

Uses real python-docx/openpyxl to build small in-memory fixture files (no
external test data files, no real Cloudinary/database — this module never
touches either). Legacy .xls (xlrd) success is exercised manually in the
conversation record with a real xlwt-authored file and real Cloudinary
round trip rather than here, to avoid adding a test-only dependency on the
unmaintained xlwt package just to author a fixture; the .xls FAILURE path
(corrupt bytes) is still covered below since it needs no fixture-authoring
library at all.
"""

import unittest
from io import BytesIO

import docx
import openpyxl
from pypdf import PdfReader

from backend.attachment_conversion import (
    MAX_EXCEL_COLUMNS_PER_SHEET_FOR_PDF,
    MAX_EXCEL_ROWS_PER_SHEET_FOR_PDF,
    convert_excel_bytes_to_pdf,
    convert_word_bytes_to_pdf,
)


def _sample_docx_bytes() -> bytes:
    document = docx.Document()
    document.add_heading("Quarterly Review", level=1)
    document.add_paragraph("First paragraph.\nWith a manual line break.")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Item"
    table.rows[0].cells[1].text = "Status"
    table.rows[1].cells[0].text = "Docs"
    table.rows[1].cells[1].text = "Done"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _sample_xlsx_bytes(sheets=None) -> bytes:
    workbook = openpyxl.Workbook()
    first = True
    for name, rows in (sheets or {"Sheet1": [["A", "B"], [1, 2]]}).items():
        ws = workbook.active if first else workbook.create_sheet(name)
        if first:
            ws.title = name
            first = False
        for row in rows:
            ws.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _pdf_page_count(pdf_bytes: bytes) -> int:
    return len(PdfReader(BytesIO(pdf_bytes)).pages)


class ConvertWordBytesToPdfTests(unittest.TestCase):
    def test_real_docx_converts_successfully(self):
        result = convert_word_bytes_to_pdf(_sample_docx_bytes(), "notes.docx")
        self.assertTrue(result.success)
        self.assertIsNone(result.error_message)
        self.assertTrue(result.pdf_bytes.startswith(b"%PDF"))
        self.assertGreaterEqual(_pdf_page_count(result.pdf_bytes), 1)

    def test_converted_pdf_contains_paragraph_and_table_text(self):
        result = convert_word_bytes_to_pdf(_sample_docx_bytes(), "notes.docx")
        text = PdfReader(BytesIO(result.pdf_bytes)).pages[0].extract_text() or ""
        self.assertIn("Quarterly Review", text)
        self.assertIn("First paragraph.", text)
        self.assertIn("Docs", text)

    def test_legacy_doc_extension_fails_with_specific_message(self):
        # Not a real OLE file — python-docx cannot parse this regardless;
        # the .doc extension selects the specific "legacy format" message.
        result = convert_word_bytes_to_pdf(b"not a real ole compound file" * 5, "old_notes.doc")
        self.assertFalse(result.success)
        self.assertIsNone(result.pdf_bytes)
        self.assertIn("Legacy .doc format", result.error_message)
        self.assertIn("ZIP download", result.error_message)

    def test_corrupt_docx_fails_with_generic_message_not_legacy_wording(self):
        result = convert_word_bytes_to_pdf(b"not a zip at all", "broken.docx")
        self.assertFalse(result.success)
        self.assertIsNone(result.pdf_bytes)
        self.assertNotIn("Legacy .doc", result.error_message)
        self.assertIn("could not be read", result.error_message)

    def test_empty_document_still_succeeds_with_placeholder_note(self):
        document = docx.Document()
        buffer = BytesIO()
        document.save(buffer)
        result = convert_word_bytes_to_pdf(buffer.getvalue(), "empty.docx")
        self.assertTrue(result.success)
        text = PdfReader(BytesIO(result.pdf_bytes)).pages[0].extract_text() or ""
        self.assertIn("no readable text content", text)


class ConvertExcelBytesToPdfTests(unittest.TestCase):
    def test_real_xlsx_converts_successfully(self):
        result = convert_excel_bytes_to_pdf(_sample_xlsx_bytes(), "book.xlsx")
        self.assertTrue(result.success)
        self.assertIsNone(result.error_message)
        self.assertTrue(result.pdf_bytes.startswith(b"%PDF"))

    def test_every_sheet_name_appears_in_the_converted_pdf(self):
        sheets = {
            "Wide Data": [["Row", "M1", "M2"], [1, 2, 3]],
            "Summary": [["Team", "Total"], ["Website", 100]],
        }
        result = convert_excel_bytes_to_pdf(_sample_xlsx_bytes(sheets), "book.xlsx")
        self.assertTrue(result.success)
        full_text = "".join(
            (page.extract_text() or "") for page in PdfReader(BytesIO(result.pdf_bytes)).pages
        )
        self.assertIn("Wide Data", full_text)
        self.assertIn("Summary", full_text)
        self.assertIn("Website", full_text)

    def test_wide_sheet_is_split_into_labelled_column_chunks(self):
        header = ["Row"] + ["Metric_" + str(i) for i in range(1, 40)]
        rows = [header, [1] + list(range(1, 40))]
        result = convert_excel_bytes_to_pdf(_sample_xlsx_bytes({"Wide": rows}), "book.xlsx")
        self.assertTrue(result.success)
        full_text = "".join(
            (page.extract_text() or "") for page in PdfReader(BytesIO(result.pdf_bytes)).pages
        )
        self.assertIn("part 1 of", full_text)
        self.assertIn("Metric_39", full_text)  # the last column is still present, in a later chunk

    def test_row_truncation_is_explicitly_noted(self):
        rows = [["Row"]] + [[i] for i in range(MAX_EXCEL_ROWS_PER_SHEET_FOR_PDF + 50)]
        result = convert_excel_bytes_to_pdf(_sample_xlsx_bytes({"Big": rows}), "book.xlsx")
        self.assertTrue(result.success)
        full_text = "".join(
            (page.extract_text() or "") for page in PdfReader(BytesIO(result.pdf_bytes)).pages
        )
        self.assertIn("Showing the first " + str(MAX_EXCEL_ROWS_PER_SHEET_FOR_PDF) + " rows", full_text)

    def test_column_truncation_is_explicitly_noted(self):
        header = ["Row"] + ["C" + str(i) for i in range(MAX_EXCEL_COLUMNS_PER_SHEET_FOR_PDF + 20)]
        rows = [header, [1] + list(range(MAX_EXCEL_COLUMNS_PER_SHEET_FOR_PDF + 20))]
        result = convert_excel_bytes_to_pdf(_sample_xlsx_bytes({"WideCols": rows}), "book.xlsx")
        self.assertTrue(result.success)
        full_text = "".join(
            (page.extract_text() or "") for page in PdfReader(BytesIO(result.pdf_bytes)).pages
        )
        self.assertIn(
            "Showing the first " + str(MAX_EXCEL_COLUMNS_PER_SHEET_FOR_PDF) + " columns", full_text
        )

    def test_corrupt_xlsx_fails_explicitly(self):
        result = convert_excel_bytes_to_pdf(b"not a zip at all", "broken.xlsx")
        self.assertFalse(result.success)
        self.assertIsNone(result.pdf_bytes)
        self.assertIn("could not be read", result.error_message)

    def test_corrupt_xls_fails_explicitly(self):
        result = convert_excel_bytes_to_pdf(b"not an ole file at all", "broken.xls")
        self.assertFalse(result.success)
        self.assertIsNone(result.pdf_bytes)
        self.assertIn("could not be read", result.error_message)

    def test_unrecognized_extension_fails_explicitly(self):
        result = convert_excel_bytes_to_pdf(b"whatever", "book.ods")
        self.assertFalse(result.success)
        self.assertIn("Unrecognized Excel file extension", result.error_message)

    def test_empty_workbook_sheet_still_succeeds_with_placeholder_note(self):
        workbook = openpyxl.Workbook()
        workbook.active.title = "Empty"
        buffer = BytesIO()
        workbook.save(buffer)
        result = convert_excel_bytes_to_pdf(buffer.getvalue(), "empty.xlsx")
        self.assertTrue(result.success)
        full_text = "".join(
            (page.extract_text() or "") for page in PdfReader(BytesIO(result.pdf_bytes)).pages
        )
        self.assertIn("This sheet is empty", full_text)


if __name__ == "__main__":
    unittest.main()
